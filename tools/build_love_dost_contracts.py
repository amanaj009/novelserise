from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_DIR = Path(r"C:\Users\abc\OneDrive\App Project Proposals")
EN_PATH = OUTPUT_DIR / "UPDATED-Love-Dost-Software-Development-Agreement-English.docx"
HI_PATH = OUTPUT_DIR / "UPDATED-Love-Dost-Software-Development-Agreement-Hindi.docx"
V2_EN_PATH = OUTPUT_DIR / "UPDATED-VERSION-2-Love-Dost-Software-Development-Agreement-English.docx"
FINAL_EN_PATH = OUTPUT_DIR / "Love-Dost-Software-Development-Agreement-English.docx"
CLIENT_APP_CONTRACT_PATH = Path(r"E:\AndroidStudioProjects\Love Dost App\Love-Dost-App-Development-Agreement-English.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def configure_document(doc, hindi=False):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    font_name = "Nirmala UI" if hindi else "Aptos"
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = font_name
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = font_name
    title._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(8)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size in (("Heading 1", 15), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = font_name
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(11)
        style.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.text = "FORVILL ENTERPRISE PVT. LTD.   |   LOVE DOST   |   UPDATED AGREEMENT"
    header.style = styles["Normal"]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor(75, 75, 75)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(90, 90, 90)


def add_title(doc, title, subtitle, language_note):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.bold = True
    r.font.size = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(language_note)
    r.italic = True
    r.font.size = Pt(9)


def add_party_table(doc, labels):
    table = doc.add_table(rows=11, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(4.65)
    values = [
        (labels[0], "FORVILL ENTERPRISES PRIVATE LIMITED"),
        (labels[1], "C5/4, 1st Floor, Pocket C5, Rohini Sector 5, New Delhi, North West Delhi, Delhi - 110085, India"),
        ("Corporate Identity Number CIN", "U62099DC2026PTC469489"),
        ("GST Identification Number GSTIN", "07AAGCF9248J1Z6"),
        ("Permanent Account Number PAN", "AAGCF9248J"),
        ("Tax Deduction Account Number TAN", "DELF13542E"),
        ("Date of Incorporation", "11 March 2026"),
        (labels[2], "Samarjeet"),
        (labels[3], "60, Nishal Faliyu, Makdum Park, Bharuch (M), P.O. Maktampur, District Bharuch, Gujarat - 392012, India"),
        (labels[4], "02 October 2026"),
        (labels[5], "Bharuch, Gujarat"),
    ]
    for i, (label, value) in enumerate(values):
        for c in table.rows[i].cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(c)
        set_cell_shading(table.cell(i, 0), "E9EFF7")
        p0 = table.cell(i, 0).paragraphs[0]
        p0.add_run(label).bold = True
        table.cell(i, 1).paragraphs[0].add_run(value)
    set_table_borders(table)
    doc.add_paragraph()


def add_clause(doc, heading, paragraphs=None, bullets=None):
    doc.add_heading(heading, level=1)
    for text in paragraphs or []:
        doc.add_paragraph(text)
    for text in bullets or []:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(text)


def add_signature_table(doc, left_title, right_title, page_break=True, left_details=None, right_details=None):
    if page_break:
        doc.add_page_break()
    doc.add_heading("Signatures" if "Provider" in left_title else "हस्ताक्षर", level=1)
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(3.3)
    table.columns[1].width = Inches(3.3)
    left_details = left_details or {}
    right_details = right_details or {}
    rows = [
        (left_title, right_title),
        (f"Name / नाम: {left_details.get('name', '__________________________')}", f"Name / नाम: {right_details.get('name', '__________________________')}"),
        (f"Designation / पद: {left_details.get('designation', '____________________')}", f"Designation / पद: {right_details.get('designation', '____________________')}"),
        ("Signature / हस्ताक्षर: ________________", "Signature / हस्ताक्षर: ________________"),
        (f"Date / दिनांक: {left_details.get('date', '________________________')}", f"Date / दिनांक: {right_details.get('date', '________________________')}"),
        (f"Place / स्थान: {left_details.get('place', '________________________')}", f"Place / स्थान: {right_details.get('place', '________________________')}"),
    ]
    for i, vals in enumerate(rows):
        for j, text in enumerate(vals):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=160, bottom=160)
            p = cell.paragraphs[0]
            run = p.add_run(text)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_shading(cell, "263B59")
    set_table_borders(table)


def build_english():
    doc = Document()
    configure_document(doc, hindi=False)
    add_title(
        doc,
        "UPDATED LOVE DOST SOFTWARE DEVELOPMENT AGREEMENT",
        "Native Android and iOS Application Development",
        "Fixed price agreement with delivery linked balance payment",
    )
    add_party_table(doc, ["Service Provider", "Provider Address", "Client", "Client Address", "Effective Date", "Court Jurisdiction"])
    doc.add_paragraph(
        "This Software Development Agreement is entered into on the Effective Date by the Service Provider and the Client identified above. It replaces the commercial payment and timeline terms stated in the earlier Love Dost proposal wherever those terms conflict with this Agreement."
    )

    add_clause(doc, "1 Purpose and Binding Terms", [
        "The Client appoints the Provider to design, develop, test, hand over and assist with deployment of the Love Dost native Android and iOS applications, together with the shared backend and administration system. The Parties intend this signed document and its signed feature addenda to be binding."
    ])

    add_clause(doc, "2 Project Scope", [
        "The base project includes the agreed Love Dost product for Android in Kotlin and iOS in Swift, connected to a shared backend and administration system. The detailed screen list, rules and acceptance checks shall be recorded in the signed feature list attached as Schedule A."
    ], [
        "User registration, login, profile, discovery, blocking and reporting.",
        "Live voice and video party rooms with host, audience, microphone and camera controls.",
        "Private voice and video calls and paid private rooms for up to three joined participants.",
        "YouTube watch-party integration, subject to YouTube availability and platform rules.",
        "Coin wallet, coin transactions, virtual gifts and transaction history.",
        "In-room games documented in the signed feature list.",
        "Reels upload, playback, feed and agreed engagement and moderation controls.",
        "One Mere Vichar written post per user per app calendar day.",
        "Ranks, points and leaderboard under approved rules.",
        "Administration tools for users, rooms, reports, coins, gifts, moderation and operational settings.",
    ])

    add_clause(doc, "3 Fixed Contract Price", [
        "The total and final development price is INR 115,000 (Rupees One Lakh Fifteen Thousand only), inclusive of applicable GST. The Provider shall not demand any additional development, professional, service, testing, handover or deployment-support fee for the signed base scope.",
        "Any tax invoice shall show the GST component included within the total amount. No verbal message, invoice or unilateral request can increase the total price. A price change is ineffective unless the Client expressly signs a separate written instrument that states the revised price; a feature addendum under Clause 7 is intended to revise features and time only and shall not by itself create an extra fee.",
        "Any optional third-party subscription, platform account charge or usage-based service shall be disclosed in advance and may be obtained only after the Client's separate written approval. Such a charge is not an additional payment to the Provider and cannot be treated as part of the Provider's development fee."
    ])

    add_clause(doc, "4 Payment Terms", [
        "The Client shall pay INR 15,000 as token money on signing this Agreement. The remaining INR 100,000 shall become due only after overall Project Completion as defined in Clause 6, including the Android and iOS deliverables.",
        "There are no interim instalments. Payment of the balance does not waive defects reported under the acceptance or maintenance provisions. The Provider shall issue the applicable receipt or GST invoice for every amount received."
    ])

    add_clause(doc, "5 Start Date and Delivery Period", [
        "The project starts on the Effective Date, meaning the date on which the last Party signs this Agreement. The Android delivery period is 100 Working Days from that date. The iOS application is allowed one additional month of development time and its delivery period is 120 Working Days from the Effective Date.",
        "A Working Day means Monday through Friday. Saturdays and Sundays are excluded. Public holidays do not automatically extend the deadline unless both Parties sign a written extension or feature addendum.",
        "Only additional days expressly stated in a signed feature addendum or written extension shall alter a deadline. An addendum must identify whether its extra days apply to Android, iOS or both. The revised Android deadline will be 100 Working Days plus its stated additional Working Days, and the revised iOS deadline will be 120 Working Days plus its stated additional Working Days."
    ])

    add_clause(doc, "6 Project Completion", [
        "Android Completion occurs when the Android Love Dost application is uploaded or deployed to the Client's Google Play Console under the Client's Play Store account or ID within the applicable 100-Working-Day deadline. Deployment to an internal testing, closed testing, open testing or production track is valid Android Completion; public listing approval is not required.",
        "iOS Completion occurs when the Provider delivers the working iOS application release build and source code, together with the agreed iOS configuration and setup information, within the applicable 120-Working-Day deadline. App Store public approval is not required unless expressly added to Schedule A.",
        "Overall Project Completion occurs only when both Android Completion and iOS Completion have occurred and the Provider has delivered the Android source code, agreed backend and admin deliverables, configuration or setup notes, and available project-specific credentials. A store review, test status or later third-party approval does not reverse timely completion once these conditions have been met."
    ])

    add_clause(doc, "7 Open Features and Time Revision", [
        "If the Client requests an open or additional feature and the Provider accepts it, the Parties shall sign a separate short feature addendum before that work begins. The addendum must describe the feature, dependencies, acceptance condition and the exact number of additional Working Days.",
        "The revised Android deadline shall be 100 Working Days plus Android-related additional Working Days, and the revised iOS deadline shall be 120 Working Days plus iOS-related additional Working Days. An unsigned chat, call, demonstration or informal request does not add a feature or extend either deadline.",
        "An accepted feature addendum does not create an additional development fee under this Agreement unless the Client separately and expressly signs a price amendment."
    ])

    add_clause(doc, "8 Delay and Refund", [
        "If the Provider does not achieve Android Completion within 100 Working Days, iOS Completion within 120 Working Days, or either validly revised deadline, the Provider shall refund every amount actually paid by the Client under this Agreement, including the token money.",
        "The complete refund must be made no later than 90 calendar days after expiry of the missed applicable deadline. The refund shall be made to the payment source or another bank account nominated by the Client in writing. No deduction, set-off or cancellation charge shall apply unless required by law.",
        "This refund clause does not apply where Android and iOS are both completed within their applicable times. Google Play deployment on a testing track is valid for Android Completion. Once timely overall completion occurs, payments are governed by Clause 4 and are not refundable merely because a store keeps an application in testing or later changes its review status."
    ])

    add_clause(doc, "9 Review Acceptance and Defects", [
        "The Provider shall share review builds and the Client shall provide consolidated feedback against the signed feature list. A defect means failure of the delivered implementation to perform an agreed function; a new or changed function is handled under Clause 7.",
        "For one year after Project Completion, the Provider shall correct reproducible defects in the delivered agreed functionality without an additional development charge. This maintenance does not include new features, redesign, content operations, third-party outages or later external API or platform-policy changes."
    ])

    add_clause(doc, "10 Client Responsibilities", [
        "The Client shall provide the branding, content, Google Play Console access, Apple Developer access where required, policy documents, credentials and timely decisions reasonably necessary for development and deployment. The Provider shall promptly identify a missing input in writing. A Client-side delay changes the contractual deadline only through a signed written extension."
    ])

    add_clause(doc, "11 Intellectual Property and Handover", [
        "Upon full payment of INR 115,000, the Client shall own the custom source code and project-specific deliverables created and paid for under this Agreement. The Provider retains ownership of pre-existing tools, general know-how and third-party components, while granting the Client the licence necessary to operate the delivered application, subject to third-party licence terms.",
        "Source repositories or archives, build instructions, relevant environment configuration and project-specific access shall be handed over no later than receipt of the final balance, with delivery and payment intended to occur concurrently."
    ])

    add_clause(doc, "12 Confidentiality Data and Security", [
        "Each Party shall keep the other Party's non-public business information, credentials and project materials confidential and use them only for this project. The Provider shall apply reasonable access controls and shall not intentionally disclose Client or user data except as required to perform the services or by law. Production account ownership should remain with the Client."
    ])

    add_clause(doc, "13 Governing Law and Courts", [
        "This Agreement is governed by the laws of India. Subject to applicable law, the competent courts at Bharuch, Gujarat shall have exclusive jurisdiction over disputes arising from or connected with this Agreement. Before filing proceedings, the Parties shall attempt in good faith to resolve the dispute through written discussion for 15 calendar days, unless urgent interim relief is required."
    ])

    add_clause(doc, "14 Notices Amendments and Electronic Signing", [
        "Formal notices and approvals shall be sent to the email addresses or verified messaging numbers stated below or later notified in writing. Amendments, feature additions, time revisions and waivers must be written and signed by both Parties.",
        "This Agreement may be signed in counterparts and by legally recognized electronic means. A scanned signed copy or valid electronic signature may be treated as an execution copy to the extent permitted by applicable law."
    ])

    add_clause(doc, "15 Entire Agreement and Priority", [
        "This Agreement, Schedule A and signed addenda contain the complete agreement for the stated project. If they conflict, a later signed amendment controls only the subject it expressly changes, followed by this Agreement, Schedule A and then the earlier Love Dost proposal. If the English and Hindi execution copies differ, the Parties shall record the agreed correction in writing; until then, the signed copy selected by both Parties as the execution copy shall control."
    ])

    doc.add_heading("Schedule A Feature and Account Confirmation", level=1)
    for item in [
        "Approved screen and feature list: [ATTACH / REFERENCE]",
        "Client Google Play Console account ID or email: [DETAILS]",
        "Client Apple Developer account ID or email: [DETAILS]",
        "Target Android and iOS versions: [DETAILS]",
        "Approved games, room limits, coin rules, Reel limits and rank rules: [DETAILS]",
        "Provider notice email and number: [DETAILS]",
        "Client notice email and number: [DETAILS]",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_signature_table(doc, "For the Provider", "For the Client")
    doc.save(EN_PATH)


def build_hindi():
    doc = Document()
    configure_document(doc, hindi=True)
    add_title(
        doc,
        "UPDATED LOVE DOST SOFTWARE DEVELOPMENT AGREEMENT",
        "नेटिव Android और iOS ऐप विकास अनुबंध",
        "निश्चित मूल्य तथा प्रोजेक्ट पूरा होने से जुड़ा शेष भुगतान",
    )
    add_party_table(doc, ["सेवा प्रदाता", "प्रदाता का पता", "ग्राहक", "ग्राहक का पता", "प्रभावी दिनांक", "न्यायालय क्षेत्राधिकार"])
    doc.add_paragraph(
        "यह सॉफ्टवेयर विकास अनुबंध प्रभावी दिनांक को ऊपर बताए गए सेवा प्रदाता और ग्राहक के बीच किया गया है। पहले Love Dost proposal की भुगतान या समय-सीमा संबंधी किसी शर्त और इस अनुबंध में अंतर होने पर इस अनुबंध की शर्त लागू होगी।"
    )

    add_clause(doc, "1 उद्देश्य और बाध्यकारी शर्तें", [
        "ग्राहक सेवा प्रदाता को Love Dost के नेटिव Android और iOS ऐप, साझा backend तथा admin system के डिजाइन, विकास, परीक्षण, handover और deployment सहायता के लिए नियुक्त करता है। यह signed अनुबंध और इसके signed feature addendum दोनों पक्षों पर बाध्यकारी होंगे।"
    ])

    add_clause(doc, "2 प्रोजेक्ट का दायरा", [
        "Base project में Kotlin में Android ऐप और Swift में iOS ऐप शामिल हैं। दोनों ऐप साझा backend और admin system से जुड़ेंगे। विस्तृत screen list, rules और acceptance checks हस्ताक्षरित feature list, Schedule A, में दर्ज होंगे।"
    ], [
        "User registration, login, profile, discovery, block और report.",
        "Host, audience, microphone और camera controls वाले live voice तथा video party rooms.",
        "Private voice/video calls और अधिकतम तीन joined participants वाला paid private room.",
        "YouTube watch party, जो YouTube की उपलब्धता और platform rules के अधीन होगी।",
        "Coin wallet, coin transactions, virtual gifts और transaction history.",
        "Signed feature list में तय किए गए in-room games.",
        "Reels upload, playback, feed, engagement और moderation controls.",
        "हर user के लिए प्रत्येक app calendar day में एक Mere Vichar post.",
        "Approved rules के अनुसार ranks, points और leaderboard.",
        "Users, rooms, reports, coins, gifts, moderation और settings के admin tools.",
    ])

    add_clause(doc, "3 निश्चित कुल मूल्य", [
        "इस प्रोजेक्ट की कुल और अंतिम development fee INR 115,000 (एक लाख पंद्रह हजार रुपये मात्र) है, जिसमें लागू GST शामिल है। Signed base scope के लिए Provider कोई अतिरिक्त development, professional, service, testing, handover या deployment-support fee नहीं मांगेगा।",
        "Tax invoice में कुल राशि के भीतर शामिल GST component दिखाया जाएगा। कोई verbal message, invoice या एकतरफा मांग कुल मूल्य नहीं बढ़ा सकती। मूल्य तभी बदलेगा जब Client अलग written price amendment पर स्पष्ट रूप से हस्ताक्षर करे। Clause 7 का feature addendum केवल feature और अतिरिक्त समय दर्ज करेगा तथा अपने आप कोई अतिरिक्त fee पैदा नहीं करेगा।",
        "किसी optional third-party subscription, platform account charge या usage-based service की जानकारी पहले लिखित रूप में देनी होगी और Client की अलग written approval के बाद ही उसे लिया जा सकेगा। ऐसा charge Provider की अतिरिक्त development fee नहीं माना जाएगा।"
    ])

    add_clause(doc, "4 भुगतान की शर्तें", [
        "Client इस अनुबंध पर हस्ताक्षर करते समय INR 15,000 token money देगा। शेष INR 100,000 केवल Clause 6 के अनुसार Android और iOS दोनों deliverables सहित overall Project Completion के बाद देय होगा।",
        "बीच में कोई installment देय नहीं होगी। शेष भुगतान करने से acceptance या maintenance के अंतर्गत defect बताने का अधिकार समाप्त नहीं होगा। Provider प्राप्त प्रत्येक payment की लागू receipt या GST invoice देगा।"
    ])

    add_clause(doc, "5 प्रारंभ दिनांक और समय सीमा", [
        "Project Effective Date से शुरू माना जाएगा। Effective Date वह दिन है जिस दिन अंतिम Party इस अनुबंध पर हस्ताक्षर करती है। Android की delivery period उस दिन से 100 Working Days होगी। iOS app के लिए एक अतिरिक्त महीने का development time रहेगा और उसकी delivery period Effective Date से 120 Working Days होगी।",
        "Working Day का अर्थ Monday से Friday है। Saturday और Sunday नहीं गिने जाएंगे। Public holiday से deadline अपने आप नहीं बढ़ेगी, जब तक दोनों Parties signed written extension या feature addendum नहीं करतीं।",
        "केवल signed feature addendum या written extension में स्पष्ट अतिरिक्त दिन deadline बदलेंगे। Addendum में लिखा होगा कि अतिरिक्त दिन Android, iOS या दोनों पर लागू हैं। Revised Android deadline 100 Working Days और उसके अतिरिक्त दिनों का कुल होगा; revised iOS deadline 120 Working Days और उसके अतिरिक्त दिनों का कुल होगा।"
    ])

    add_clause(doc, "6 Project Completion कब माना जाएगा", [
        "Android Completion तब होगा जब Android Love Dost app लागू 100-Working-Day deadline के भीतर Client के Google Play Console में Client के Play Store account या ID पर upload या deploy कर दिया जाए। Internal testing, closed testing, open testing या production track का deployment मान्य होगा; public listing approval आवश्यक नहीं है।",
        "iOS Completion तब होगा जब Provider लागू 120-Working-Day deadline के भीतर working iOS release build और source code, साथ में agreed iOS configuration तथा setup information, Client को दे देता है। App Store की public approval आवश्यक नहीं होगी, जब तक Schedule A में स्पष्ट रूप से न जोड़ी जाए।",
        "Overall Project Completion तभी होगा जब Android Completion और iOS Completion दोनों हो जाएं तथा Android source code, agreed backend और admin deliverables, configuration या setup notes तथा project-specific उपलब्ध credentials भी दे दिए जाएं। Store review, test status या बाद की third-party approval timely completion को वापस नहीं बदलेगी।"
    ])

    add_clause(doc, "7 Open Features और अतिरिक्त दिन", [
        "यदि Client कोई open या additional feature मांगता है और Provider उसे स्वीकार करता है, तो उस काम के शुरू होने से पहले दोनों Parties अलग छोटा feature addendum sign करेंगी। Addendum में feature, dependencies, acceptance condition और अतिरिक्त Working Days की सही संख्या लिखी जाएगी।",
        "Revised Android deadline 100 Working Days और Android से जुड़े अतिरिक्त दिनों का कुल होगा। Revised iOS deadline 120 Working Days और iOS से जुड़े अतिरिक्त दिनों का कुल होगा। Unsigned chat, call, demo या informal request से feature add या कोई deadline extend नहीं होगी।",
        "Accepted feature addendum से इस अनुबंध के अंतर्गत कोई अतिरिक्त development fee नहीं बनेगी, जब तक Client अलग और स्पष्ट price amendment पर हस्ताक्षर न करे।"
    ])

    add_clause(doc, "8 देरी और Refund", [
        "यदि Provider Android Completion 100 Working Days में, iOS Completion 120 Working Days में, या इनमें से किसी की validly revised deadline में पूरा नहीं करता, तो Provider Client द्वारा इस अनुबंध के अंतर्गत वास्तव में दी गई हर राशि, token money सहित, वापस करेगा।",
        "पूरा refund miss हुई लागू deadline समाप्त होने के बाद अधिकतम 90 calendar days के भीतर करना होगा। Refund उसी payment source या Client द्वारा लिखित रूप में बताए गए bank account में होगा। कानून द्वारा आवश्यक न होने पर कोई deduction, set-off या cancellation charge लागू नहीं होगा।",
        "यदि Android और iOS दोनों अपनी लागू समय सीमाओं में पूरे हो जाते हैं, तो यह refund clause लागू नहीं होगा। Google Play का testing-track deployment Android Completion के लिए मान्य है। Timely overall completion के बाद payment Clause 4 के अनुसार होगा और केवल store testing status या बाद में बदली review status के कारण refundable नहीं होगा।"
    ])

    add_clause(doc, "9 Review Acceptance और Defects", [
        "Provider review builds देगा और Client signed feature list के अनुसार consolidated feedback देगा। Defect का अर्थ agreed function का delivered implementation में काम न करना है। नया या बदला हुआ function Clause 7 के अनुसार संभाला जाएगा।",
        "Project Completion के बाद एक वर्ष तक Provider delivered agreed functionality के reproducible defects बिना अतिरिक्त development charge के ठीक करेगा। इसमें नए features, redesign, content operations, third-party outage या बाद में बदली external API/platform policy शामिल नहीं है।"
    ])

    add_clause(doc, "10 Client की जिम्मेदारियां", [
        "Client development और deployment के लिए reasonably आवश्यक branding, content, Google Play Console access, आवश्यक होने पर Apple Developer access, policy documents, credentials और समय पर decisions देगा। Provider किसी missing input की जानकारी तुरंत लिखित रूप में देगा। Client-side delay से contractual deadline केवल signed written extension द्वारा बदलेगी।"
    ])

    add_clause(doc, "11 Intellectual Property और Handover", [
        "कुल INR 115,000 का पूरा payment मिलने पर इस अनुबंध के अंतर्गत बनाए और pay किए गए custom source code तथा project-specific deliverables का ownership Client को मिलेगा। Provider अपने pre-existing tools, सामान्य know-how और third-party components का ownership रखेगा, लेकिन Client को delivered app चलाने के लिए आवश्यक licence देगा, जो third-party licence terms के अधीन होगा।",
        "Source repositories या archives, build instructions, relevant environment configuration और project-specific access अंतिम balance मिलने तक handover किए जाएंगे। Delivery और final payment एक साथ होने का उद्देश्य है।"
    ])

    add_clause(doc, "12 Confidentiality Data और Security", [
        "दोनों Parties दूसरी Party की non-public business information, credentials और project material को confidential रखेंगी और केवल इस project के लिए उपयोग करेंगी। Provider reasonable access controls रखेगा और services या कानून की आवश्यकता के अलावा Client या user data को जानबूझकर disclose नहीं करेगा। Production accounts Client के ownership में रहने चाहिए।"
    ])

    add_clause(doc, "13 लागू कानून और न्यायालय", [
        "यह अनुबंध भारत के कानूनों द्वारा नियंत्रित होगा। लागू कानून के अधीन, इस अनुबंध से संबंधित विवादों पर Bharuch, Gujarat के सक्षम न्यायालयों का exclusive jurisdiction होगा। Court proceeding शुरू करने से पहले Parties 15 calendar days तक written discussion द्वारा समाधान का सद्भावनापूर्वक प्रयास करेंगी, सिवाय urgent interim relief की आवश्यकता के।"
    ])

    add_clause(doc, "14 Notices Amendments और Electronic Signing", [
        "Formal notices और approvals नीचे दिए email addresses या verified messaging numbers पर भेजे जाएंगे। Amendment, feature addition, time revision और waiver लिखित और दोनों Parties द्वारा signed होना चाहिए।",
        "यह अनुबंध counterparts और कानून द्वारा मान्य electronic means से sign किया जा सकता है। लागू कानून की अनुमति के अनुसार scanned signed copy या valid electronic signature को execution copy माना जा सकता है।"
    ])

    add_clause(doc, "15 पूरा अनुबंध और प्राथमिकता", [
        "यह Agreement, Schedule A और signed addenda stated project का पूरा agreement हैं। Conflict होने पर बाद का signed amendment केवल स्पष्ट रूप से बदले subject पर लागू होगा, फिर यह Agreement, Schedule A और उसके बाद पुराना Love Dost proposal लागू होगा। English और Hindi execution copies में अंतर होने पर Parties agreed correction लिखित रूप में दर्ज करेंगी; तब तक दोनों द्वारा execution copy के रूप में चुनी signed copy लागू होगी।"
    ])

    doc.add_page_break()
    doc.add_heading("Schedule A Feature और Account Confirmation", level=1)
    for item in [
        "Approved screen और feature list: [ATTACH / REFERENCE]",
        "Client Google Play Console account ID या email: [DETAILS]",
        "Client Apple Developer account ID या email: [DETAILS]",
        "Target Android और iOS versions: [DETAILS]",
        "Approved games, room limits, coin rules, Reel limits और rank rules: [DETAILS]",
        "Provider notice email और number: [DETAILS]",
        "Client notice email और number: [DETAILS]",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_signature_table(doc, "सेवा प्रदाता की ओर से", "ग्राहक की ओर से", page_break=False)
    doc.save(HI_PATH)


def add_v2_price_table(
    doc,
    token_amount="INR 15,000",
    android_balance="INR 60,000",
    android_before_gst="INR 63,559.32",
    android_gst="INR 11,440.68",
    android_price="INR 75,000",
    ios_before_gst="INR 33,898.31",
    ios_gst="INR 6,101.69",
    ios_price="INR 40,000",
    total_before_gst="INR 97,457.63",
    total_gst="INR 17,542.37",
    total_price="INR 115,000",
):
    table = doc.add_table(rows=4, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [2.05, 1.15, 1.15, 1.15, 1.25]
    for idx, width in enumerate(widths):
        table.columns[idx].width = Inches(width)
    headers = ["Project", "Before GST", "GST 18%", "GST Inclusive", "Payment Trigger"]
    data = [
        ["Android app plus shared backend and admin", android_before_gst, android_gst, android_price, f"{token_amount} token plus {android_balance} on Android Completion"],
        ["iOS app and shared-system integration", ios_before_gst, ios_gst, ios_price, f"{ios_price} on iOS Completion"],
        ["Total", total_before_gst, total_gst, total_price, "No cross-hold between projects"],
    ]
    for j, text in enumerate(headers):
        cell = table.cell(0, j)
        set_cell_shading(cell, "263B59")
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    for i, row in enumerate(data, start=1):
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            if i % 2 == 0:
                set_cell_shading(cell, "F3F6FA")
            cell.paragraphs[0].add_run(text)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=100, start=100, bottom=100, end=100)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8.5)
    set_table_borders(table)
    doc.add_paragraph()


def build_v2_english():
    doc = Document()
    configure_document(doc, hindi=False)
    add_title(
        doc,
        "UPDATED VERSION 2 LOVE DOST SOFTWARE DEVELOPMENT AGREEMENT",
        "Separate Android and iOS Projects",
        "Independent completion payment and refund obligations",
    )
    add_party_table(doc, ["Service Provider", "Provider Address", "Client", "Client Address", "Effective Date", "Court Jurisdiction"])
    doc.add_paragraph(
        "This Software Development Agreement is entered into on the Effective Date by the Service Provider and the Client identified above. It replaces the commercial payment, completion, refund and timeline terms in the earlier Love Dost proposal and the prior updated agreement wherever those terms conflict with this Version 2 Agreement."
    )

    add_clause(doc, "1 Agreement Structure", [
        "The work is divided into two independently priced and independently completed projects: the Android Project and the iOS Project. Completion, payment, delay and refund shall be assessed separately for each project. Completion or delay of one project shall not postpone payment for the other completed project."
    ])

    add_clause(doc, "2 Android Project", [
        "The Android Project includes the native Kotlin Android application, the shared backend, and the administration system required for the signed Android feature list. It includes agreed account and profile functions, discovery and moderation, live voice and video parties, private calls and rooms, YouTube watch party, wallet and gifts, games, Reels, Mere Vichar, ranks, leaderboard and administration functions. The signed feature list in Schedule A controls the measurable Android scope."
    ])

    add_clause(doc, "3 iOS Project", [
        "The iOS Project includes the native Swift iOS application and its integration with the agreed shared backend and administration system. It shall provide the corresponding signed iOS features, subject to iOS platform structure, Apple requirements and the acceptance checks in Schedule A. The Parties acknowledge that native iOS builds and Apple processes can require different implementation work and may take longer than Android."
    ])

    add_clause(doc, "4 Fixed Price and GST", [
        "The combined fixed price for both projects is INR 115,000 (Rupees One Lakh Fifteen Thousand only), inclusive of applicable GST. The Android Project price is INR 75,000 inclusive of GST. The iOS Project price is INR 40,000 inclusive of GST. The Provider shall not demand any additional development, testing, handover or deployment-support fee for the signed scope."
    ])
    add_v2_price_table(doc)

    add_clause(doc, "5 Android Payment", [
        "The Client shall pay INR 15,000 as token money on signing this Agreement. The complete token amount is allocated to the Android Project price.",
        "After Android Completion, the remaining Android balance of INR 60,000 shall become due and payable against the Provider's applicable GST invoice. The total amount payable for the Android Project shall not exceed INR 75,000 inclusive of GST.",
        "Android payment is independent. The Client shall not withhold or delay the INR 60,000 Android balance because the iOS Project remains under development, requires a different technical structure, encounters an Apple-related delay, or is not yet complete."
    ])

    add_clause(doc, "6 iOS Payment", [
        "The complete iOS Project price is INR 40,000 inclusive of GST. It shall become due and payable only after iOS Completion against the Provider's applicable GST invoice.",
        "No iOS payment is due merely because the Android Project has been completed. An Android payment does not count as an iOS payment, and an iOS delay shall not affect the completed Android Project or its payment."
    ])

    add_clause(doc, "7 Third Party Costs", [
        "All third-party costs are the Client's responsibility and shall be paid by the Client. These costs may include domain registration, hosting, cloud servers, databases, storage, CDN, voice or video services, SMS or OTP, payment-gateway charges, Google Play fees, Apple Developer fees, third-party APIs, purchased code, licences and usage-based service charges.",
        "The Provider shall disclose a proposed third-party cost before activation and obtain the Client's written approval. The Client may pay the third-party provider directly or reimburse the exact approved amount against evidence. The Provider shall not add a markup or treat a third-party cost as an additional development fee. No unapproved third-party cost shall be payable by the Client."
    ])

    add_clause(doc, "8 Start Date and Separate Deadlines", [
        "Both projects start on the Effective Date, meaning the date on which the last Party signs this Agreement. A Working Day means Monday through Friday; Saturdays and Sundays are excluded. A public holiday does not automatically extend a deadline unless both Parties sign a written extension.",
        "The Android Project deadline is 100 Working Days from the Effective Date. The iOS Project deadline is 120 Working Days from the Effective Date. The additional iOS time recognizes its separate native structure, build process and Apple-related implementation risk."
    ])

    add_clause(doc, "9 Separate Completion Events", [
        "Android Completion occurs when the Android application satisfies the signed Android acceptance checks, is uploaded or deployed to the Client's Google Play Console under the Client's account or ID, and the Android source code plus agreed shared backend and admin deliverables are made available. Internal, closed, open or production testing track deployment is valid; public listing approval is not required.",
        "iOS Completion occurs when the iOS application satisfies the signed iOS acceptance checks, its release build and source code are delivered, and it is uploaded to the Client's App Store Connect or TestFlight account where the Client has supplied the required access. TestFlight or another Apple testing status is valid; public App Store approval is not required.",
        "Each project is complete independently when its own conditions are met. There is no requirement to wait for both projects before recognizing or paying for a completed project."
    ])

    add_clause(doc, "10 Open Features and Revised Days", [
        "Any accepted open or additional feature shall be recorded before work begins in a separate short addendum signed by both Parties. The addendum shall identify whether the feature belongs to Android, iOS or both, describe its acceptance condition and state the exact additional Working Days for each affected project.",
        "The revised Android deadline will be 100 Working Days plus Android additional days. The revised iOS deadline will be 120 Working Days plus iOS additional days. A change to one project does not extend the other project unless the signed addendum expressly says so. An accepted feature addendum does not create an additional development fee unless the Client separately signs an express price amendment."
    ])

    add_clause(doc, "11 Project Specific Delay and Refund", [
        "If Android Completion is not achieved within its applicable original or revised deadline, the Provider shall refund only the amounts actually received and allocated to the Android Project, including the INR 15,000 token money. The refund shall be completed within 90 calendar days after the missed Android deadline.",
        "If iOS Completion is not achieved within its applicable original or revised deadline, the Provider shall refund only the amounts actually received and allocated to the iOS Project. If no iOS amount has been received, no monetary refund is due, although the Client has no obligation to pay the INR 40,000 iOS price. Any iOS refund due shall be completed within 90 calendar days after the missed iOS deadline.",
        "Failure or refund relating to one project does not cancel, reverse or make refundable the fee for the other project if that other project was completed on time. Google Play or TestFlight testing status is valid completion when the applicable conditions in Clause 9 are met."
    ])

    add_clause(doc, "12 Review Defects and Maintenance", [
        "The Client shall review each project against its signed feature list and give consolidated feedback. A defect is a failure of delivered implementation to perform an agreed function; a new or changed function is handled under Clause 10.",
        "For one year after each project's respective Completion date, the Provider shall correct reproducible defects in that project's delivered agreed functionality without an additional development charge. Maintenance does not include new features, redesign, content operations, third-party outages or later external API or platform-policy changes."
    ])

    add_clause(doc, "13 Client Responsibilities", [
        "The Client shall provide branding, content, Google Play Console access, Apple Developer and App Store Connect access, policy documents, credentials, approved third-party accounts and timely decisions reasonably necessary for the relevant project. The Provider shall identify missing inputs in writing. A Client-side delay changes a deadline only through a signed written extension."
    ])

    add_clause(doc, "14 Intellectual Property and Separate Handover", [
        "Upon full payment of the Android Project price, the Client shall own the custom Android source code and project-specific shared backend and admin deliverables created for that project. Upon full payment of the iOS Project price, the Client shall own the custom iOS source code and iOS-specific deliverables.",
        "The Provider retains pre-existing tools, general know-how and third-party components while granting the licences necessary to operate the delivered applications, subject to third-party licence terms. Each project's handover and ownership transfer shall be handled independently."
    ])

    add_clause(doc, "15 Confidentiality Data and Security", [
        "Each Party shall keep the other Party's non-public business information, credentials and project materials confidential and use them only for the projects. The Provider shall apply reasonable access controls and shall not intentionally disclose Client or user data except as required to perform the services or by law. Production accounts should remain under the Client's ownership."
    ])

    add_clause(doc, "16 Governing Law and Courts", [
        "This Agreement is governed by the laws of India. Subject to applicable law, the competent courts at Bharuch, Gujarat shall have exclusive jurisdiction over disputes connected with this Agreement. Before filing proceedings, the Parties shall attempt in good faith to resolve the dispute through written discussion for 15 calendar days, unless urgent interim relief is required."
    ])

    add_clause(doc, "17 Notices Amendments and Electronic Signing", [
        "Formal notices and approvals shall be sent to the notified email addresses or verified messaging numbers. Amendments, feature additions, price changes, time revisions and waivers must be written and signed by both Parties.",
        "This Agreement may be signed in counterparts and by legally recognized electronic means. A scanned signed copy or valid electronic signature may be treated as an execution copy to the extent permitted by applicable law."
    ])

    add_clause(doc, "18 Entire Agreement and Severability", [
        "This Version 2 Agreement, Schedule A and signed addenda contain the complete agreement for the two projects. A later signed amendment controls only the subject it expressly changes. The independent Android and iOS structure is severable: invalidity, delay, refund or termination affecting one project does not automatically invalidate the completed obligations for the other project."
    ])

    doc.add_page_break()
    doc.add_heading("Schedule A Separate Acceptance and Account Confirmation", level=1)
    for item in [
        "Approved Android screen, feature and acceptance list: [ATTACH / REFERENCE]",
        "Approved iOS screen, feature and acceptance list: [ATTACH / REFERENCE]",
        "Client Google Play Console account ID or email: [DETAILS]",
        "Client Apple Developer or App Store Connect account ID or email: [DETAILS]",
        "Target Android versions and devices: [DETAILS]",
        "Target iOS versions and devices: [DETAILS]",
        "Approved games, room limits, coin rules, Reel limits and rank rules: [DETAILS]",
        "Provider notice email and number: [DETAILS]",
        "Client notice email and number: [DETAILS]",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    add_signature_table(doc, "For the Provider", "For the Client", page_break=False)
    doc.save(V2_EN_PATH)


def build_main_english():
    doc = Document()
    configure_document(doc, hindi=False)
    doc.settings.odd_and_even_pages_header_footer = False
    doc.sections[0].top_margin = Inches(0.82)
    main_header = doc.sections[0].header.paragraphs[0]
    main_header.text = ""
    main_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in main_header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(7.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(75, 75, 75)
    add_title(
        doc,
        "LOVE DOST APP DEVELOPMENT AGREEMENT",
        "Native Android and iOS Applications with Shared Backend and Administration",
        "Integrated project scope commercial terms completion rules and handover",
    )
    add_party_table(doc, ["Service Provider", "Provider Address", "Client", "Client Address", "Effective Date", "Court Jurisdiction"])
    doc.add_paragraph(
        "This App Development Agreement is entered into on the Effective Date by the Service Provider and the Client identified above. It contains the complete commercial and delivery terms for the Love Dost Android and iOS applications and supersedes all earlier proposals, quotations and unsigned draft agreements concerning the same subject."
    )

    add_clause(doc, "1 Purpose and Appointment", [
        "The Client appoints the Provider to design, develop, integrate, test, hand over and assist with deployment of the Love Dost social entertainment platform. The engagement includes a native Android application, a native iOS application, a shared backend and an administration system, subject to the detailed scope and acceptance conditions in this Agreement."
    ])

    add_clause(doc, "2 Independent Android and iOS Projects", [
        "The work is divided into two independently priced and independently completed projects: the Android Project and the iOS Project. Completion, payment, delay, handover and refund shall be assessed separately for each project. Completion or delay of one project shall not postpone payment for the other completed project.",
        "The Android Project includes the native Kotlin Android application and the initial shared backend and administration system. The iOS Project includes the native Swift iOS application and its integration with the agreed shared backend and administration system. Both applications shall provide the agreed core user experience, subject to platform-specific implementation differences."
    ])

    add_clause(doc, "3 Complete Product Scope", [
        "The following functions form the intended base scope for Android and iOS. The Parties shall record measurable screen-level acceptance checks, limits and configuration choices in Schedule A before detailed implementation."
    ])

    doc.add_heading("3.1 Accounts Profiles and Discovery", level=2)
    for text in [
        "User registration and login through the approved account method.",
        "User profile with display name, photograph, introduction and approved public details.",
        "Discovery of users, live rooms and relevant social content.",
        "Basic social interactions together with user blocking and reporting controls.",
        "Search, privacy and account settings stated in Schedule A.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("3.2 Live Voice and Video Parties", level=2)
    for text in [
        "Host-created live voice and video party rooms with host, participant and audience roles.",
        "Host controls to open or close a room and control microphone, speaker, camera and participation access.",
        "Audience joining, microphone requests, participant list, agreed room chat and gift activity.",
        "Room reporting, moderation and removal controls.",
        "Room capacity and video-quality targets based on the approved real-time service tier.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("3.3 Private Calls and Paid Private Rooms", level=2)
    for text in [
        "Private one-to-one voice and video calling.",
        "Paid private video rooms supporting a maximum of three joined participants.",
        "Visible coin price before entry, wallet-balance verification, coin deduction and transaction-history entry.",
        "Rules for pricing authority, host earnings if any, rejoining, invite or discovery mode and interrupted sessions as approved in Schedule A.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("3.4 YouTube Watch Party", level=2)
    for text in [
        "A host may select a permitted embeddable YouTube video for synchronized viewing with room voice conversation.",
        "Approved controls for play, pause, seek, late joining, unavailable videos and host departure.",
        "Availability remains subject to YouTube embedding permissions, region restrictions, network conditions and platform rules.",
        "The application shall not sell access to a YouTube video; any eligible charge relates only to the app service or room experience.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_page_break()
    doc.add_heading("3.5 Coins Gifts and Games", level=2)
    for text in [
        "Coin wallet, approved coin-purchase flow, current balance and transaction history.",
        "Configurable virtual gift catalogue with gift names, visuals and coin values.",
        "Gift sending in eligible rooms and agreed transaction review through the administration system.",
        "A clearly labelled Games button or menu within each eligible voice or video party room, placed in the room controls or bottom action toolbar as finalized in the signed user-interface design.",
        "The Games control shall open the catalogue of Client-approved licensed third-party games, allow the host to launch an available game, allow eligible room participants to join and display the result or score returned by the approved game service.",
        "The Provider shall install and integrate the approved third-party game SDK, module, web view or deep-link launch method. Original game development, game source code, artwork, licences, game servers and vendor charges are not included in the fixed development price.",
        "Lucky Bag and PK War are reserved future feature concepts described in Schedule A. They do not form part of base acceptance until their final rules and additional time are recorded in a signed feature addendum under Clause 11.",
        "Store billing and payment flows shall follow the applicable Google and Apple platform rules.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("3.6 Reels and Mere Vichar", level=2)
    for text in [
        "Short-video Reel upload, vertical scrolling feed, playback and approved basic engagement such as like, comment and report.",
        "Reel duration, size, format, feed order, storage limits and moderation behavior shall be stated in Schedule A.",
        "Mere Vichar shall allow one written thought per user per app calendar day and prevent a second post for the same day.",
        "The daily thought shall appear in the approved public feed and profile view, subject to agreed text length, timezone, edit, delete, reaction and comment rules.",
        "Reporting and administration moderation shall apply to Reels and Mere Vichar content.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("3.7 Ranks and Leaderboard", level=2)
    for text in [
        "User level or rank displayed on the profile and an in-app leaderboard.",
        "Points based on approved activities such as participation, engagement, gifts or game results.",
        "Schedule A shall define point weights, rank names, reset periods, tie rules and anti-abuse checks.",
        "Administrators may review suspicious activity and apply corrections under approved rules. Rank alone creates no cash entitlement.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("3.8 Administration and Community Operations", level=2)
    for text in [
        "Administration tools for users, rooms, reports, gifts, coins, ranks, content and operational settings.",
        "Approved moderation actions, including disabling a user account or room and reviewing reported activity.",
        "Permission levels, report categories, available audit information and data-retention requirements stated in Schedule A.",
        "After launch, the Client is responsible for daily moderation, community-policy decisions, content operations and user support unless separately agreed.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    add_clause(doc, "4 Technology and Quality Approach", [
        "The Android application shall be developed natively in Kotlin and the iOS application natively in Swift. Both shall connect to a shared backend and administration system. Platform SDKs or an embedded web view may be used where an approved integration requires them, while the principal user experience remains native.",
        "The Provider shall implement reasonable session and wallet-state handling, error messages, access checks and technical logging needed to investigate faults. Existing or purchased source code may be used only after appropriate licence, security, quality and compatibility review and remains subject to its licence terms.",
        "Testing shall cover the agreed critical journeys on representative devices, including sign-in, room creation and joining, microphone and camera controls, private calls, paid entry, coins, gifts, watch-party playback, games, Reel upload and playback, the daily Mere Vichar limit, ranks, reporting and administration actions."
    ])

    add_clause(doc, "5 Fixed Price and GST", [
        "The combined fixed price for both projects is INR 120,000 (Rupees One Lakh Twenty Thousand only), inclusive of applicable GST. The Android Project price is INR 70,000 inclusive of GST. The iOS Project price is INR 50,000 inclusive of GST.",
        "The Provider shall not demand an additional development, professional, testing, handover or deployment-support fee for the signed base scope or an accepted open feature documented under Clause 11. A separate product or service outside Love Dost may be declined or proposed under a separate independent agreement."
    ])
    add_v2_price_table(
        doc,
        token_amount="INR 20,000",
        android_balance="INR 50,000",
        android_before_gst="INR 59,322.03",
        android_gst="INR 10,677.97",
        android_price="INR 70,000",
        ios_before_gst="INR 42,372.88",
        ios_gst="INR 7,627.12",
        ios_price="INR 50,000",
        total_before_gst="INR 101,694.91",
        total_gst="INR 18,305.09",
        total_price="INR 120,000",
    )

    add_clause(doc, "6 Android Payment", [
        "The Client shall pay INR 20,000 as token money on signing this Agreement. The entire token amount is allocated to the Android Project price.",
        "After Android Completion, the remaining Android balance of INR 50,000 shall become due and payable against the applicable GST invoice. The total Android Project liability shall not exceed INR 70,000 inclusive of GST.",
        "The Client shall not withhold or delay the INR 50,000 Android balance because the iOS Project remains under development, requires a different technical structure, encounters an Apple-related delay or is not yet complete."
    ])

    add_clause(doc, "7 iOS Payment", [
        "The iOS Project price is INR 50,000 inclusive of GST and becomes due and payable only after iOS Completion against the applicable GST invoice. No iOS payment is due merely because Android is complete. Android payment does not count as iOS payment, and an iOS delay does not affect payment for a completed Android Project."
    ])

    add_clause(doc, "8 Third Party Costs", [
        "All third-party costs are the Client's responsibility and shall be paid by the Client. These may include domain registration, hosting, cloud servers, databases, storage, CDN, real-time voice or video services, SMS or OTP, payment-gateway charges, Google Play fees, Apple Developer fees, third-party APIs, purchased code, licences and usage-based charges.",
        "The Client shall procure or provide, own and fund the production server, cloud, database, storage and related infrastructure accounts. The Provider shall recommend the reasonable technical configuration and shall configure or deploy the agreed backend using the access supplied by the Client. The Provider is not required to purchase, finance or permanently hold a production infrastructure account in its own name.",
        "Where the approved provider offers a suitable India region that supports the agreed functions, the production backend, primary database, user-upload storage and available backups shall use an India region as the preferred default. If a required service or feature is not reasonably available in an India region, another region may be used only after the Provider explains the reason and the Client gives written approval. The selected region remains subject to the approved vendor's availability, architecture and service terms.",
        "Before activating a paid third-party service, the Provider shall identify the proposed vendor, pricing basis and a reasonable monthly, annual or usage-based cost estimate based on the agreed usage assumptions. The Provider shall obtain the Client's written approval for the service and its approved budget or spending cap. Where the vendor supports spending alerts or service limits, the Provider shall configure the available controls requested by the Client.",
        "The Client may pay the third-party provider directly or reimburse the exact approved amount against evidence. The Provider shall not add a markup, treat a third-party cost as an additional development fee or knowingly commit the Client above an approved cap without further written approval. Vendor price changes, taxes and actual usage may differ from an estimate and remain payable by the Client only to the extent validly incurred under the approved account or plan. No unapproved third-party cost is payable by the Client."
    ])

    add_clause(doc, "9 Start Date and Calendar Days", [
        "Both projects start on the Effective Date, being the date on which the last Party signs this Agreement. Each Party shall write or electronically record its actual signing date in the signature section. If the Parties sign on different dates, the later signature date is the Effective Date even if the Effective Date field on the first page is incomplete. All periods stated in days are calendar days unless this Agreement expressly states otherwise."
    ])

    add_clause(doc, "10 Separate Delivery Deadlines", [
        "The Android Project deadline is 90 calendar days from the Effective Date. The iOS Project deadline is 150 calendar days from the Effective Date, subject to the Client-delay exclusions and signed feature extensions stated below.",
        "If the Provider requests in writing any Client input, access, account, credential, approval, content or decision reasonably required to continue the affected work, and the Client does not provide a complete and usable response by the requested reasonable date, the resulting blocked period shall not count toward the applicable 90-day or 150-day deadline. The affected deadline automatically extends by the same number of calendar days without requiring a separate amendment.",
        "Client-delay exclusions include delay in providing or activating server or hosting details, domain or cloud access, API credentials, payment-gateway access, Google Play Console account or ID, Apple Developer or App Store Connect account or ID, store permissions, branding, content, legal policies, moderation rules, feature decisions, testing feedback or approvals.",
        "The excluded period begins on the day after the reasonable response date stated in the Provider's written request, or on the date the affected work actually becomes blocked, whichever is later. It ends when the required item is supplied in complete and usable form. The Provider shall notify the Client in writing that the affected work is blocked and shall confirm when work can resume. Only the affected project's deadline is extended unless the missing item blocks both projects.",
        "Except where the missing item immediately blocks work, creates a security risk or is subject to an external expiry, the Provider shall send the initial written request and at least one written follow-up reminder before treating later days as excluded Client-delay days. A follow-up reminder is unnecessary where the Client has already confirmed that the item will not be supplied by the requested date or the work is objectively blocked.",
        "A late response does not by itself extend a deadline. Only the actual documented period during which the missing Client item prevents the affected work from reasonably continuing is excluded. The Provider shall continue unaffected work where reasonably possible, take reasonable steps to reduce the delay and shall not count the same blocked day more than once."
    ])

    add_clause(doc, "11 Open Features and Signed Addenda", [
        "If the Client requests an open or additional Love Dost feature and the Provider accepts it, the Parties shall sign a separate short feature addendum before the work begins. The addendum shall identify whether the feature affects Android, iOS or both, describe the feature and dependencies, state its acceptance condition and specify the exact additional calendar days for each affected project.",
        "The revised Android deadline will be 90 calendar days plus the additional days stated in its signed addendum and any Client-delay days excluded under Clause 10. The revised iOS deadline will be 150 calendar days plus the corresponding additional and excluded days. A change to one project does not extend the other unless the addendum or blocking event expressly affects both. An unsigned chat, call, demonstration or informal request does not change the scope or deadline.",
        "An accepted Love Dost feature addendum carries no additional development fee. New third-party costs remain governed by Clause 8."
    ])

    add_clause(doc, "12 Android Completion", [
        "Android Completion occurs when the Android application satisfies the signed Android acceptance checks, is uploaded or deployed to the Client's Google Play Console under the Client's account or ID, and the Android source code plus the agreed shared backend, administration, configuration and setup deliverables are made available.",
        "Upload or deployment by itself does not constitute Android Completion. Deployment to an internal, closed, open or production testing track satisfies only the store-deployment element of Completion. Android Completion occurs only when all requirements in the preceding paragraph, including the signed Android acceptance checks, are satisfied. Public listing approval is not required. Later Google review timing or testing status does not reverse a timely Android Completion.",
        "The period beginning when the Provider successfully submits or uploads the eligible Android build to the Client's Google Play Console and ending when Google completes its review, requests further information, approves, rejects or publishes the app is the Google Review Period. The duration of the Google Review Period is controlled solely by Google and is excluded from the Android Project deadline. Those review days do not constitute Provider delay and do not trigger a refund or delay penalty. The date on which the app becomes publicly live remains subject to Google's review time, policies, systems and decision.",
        "If Google rejects the submission solely because the Provider's delivered build fails the signed Android acceptance checks or the Provider's submission package materially fails Google's published technical submission requirements in force on the submission date, the Provider shall correct and resubmit the affected item without an additional development charge. Time spent by the Provider making that correction is not part of the excluded Google Review Period; a new Google Review Period begins upon successful resubmission."
    ])

    add_clause(doc, "13 iOS Completion", [
        "iOS Completion occurs when the iOS application satisfies the signed iOS acceptance checks, its release build and source code are delivered, and it is uploaded to the Client's App Store Connect or TestFlight account where the Client has provided the required access.",
        "Upload or TestFlight availability by itself does not constitute iOS Completion. TestFlight or another Apple testing status satisfies only the store-deployment element of Completion. iOS Completion occurs only when all requirements in the preceding paragraph, including the signed iOS acceptance checks and delivery of the release build and source code, are satisfied. Public App Store approval is not required. Each project is complete independently; neither Party must wait for both projects before recognizing and paying for a completed project.",
        "Each application shall be submitted and launched separately as soon as its own Completion conditions are satisfied and the Client has provided the required store account access. The Provider shall not delay the Android submission or launch process until iOS Completion, and shall not delay the iOS submission or launch process until Android Completion. Each application's public-live date remains subject to the applicable store's review, approval, policies and processing time.",
        "The period beginning when the Provider successfully submits or uploads the eligible iOS build to the Client's App Store Connect or TestFlight account and ending when Apple completes its review, requests further information, approves, rejects or publishes the app is the Apple Review Period. The duration of the Apple Review Period is controlled solely by Apple and is excluded from the iOS Project deadline. Those review days do not constitute Provider delay and do not trigger a refund or delay penalty. The date on which the app becomes publicly live remains subject to Apple's review time, policies, systems and decision.",
        "If Apple rejects the submission solely because the Provider's delivered build fails the signed iOS acceptance checks or the Provider's submission package materially fails Apple's published technical submission requirements in force on the submission date, the Provider shall correct and resubmit the affected item without an additional development charge. Time spent by the Provider making that correction is not part of the excluded Apple Review Period; a new Apple Review Period begins upon successful resubmission."
    ])

    add_clause(doc, "14 Review and Acceptance", [
        "The Provider shall share relevant builds or administration demonstrations and a list of completed items. The Client shall review them against the applicable signed acceptance checks and provide one consolidated set of observations through its nominated contact.",
        "A defect means delivered behavior that fails the signed feature list or acceptance check. A request that changes approved behavior, adds a screen, rule, game or integration, or increases an agreed limit is an additional feature governed by Clause 11.",
        "A store's approval time, an external service outage or a later platform-policy change is not a defect in the Provider's delivered implementation."
    ])

    add_clause(doc, "15 Project Specific Delay and Refund", [
        "If Android Completion is not achieved within its applicable deadline after adding signed feature-extension days and excluding documented Client-delay days, the Provider shall refund only amounts actually received and allocated to the Android Project, including the INR 20,000 token money. The refund shall be completed within 30 calendar days after the missed Android deadline.",
        "If iOS Completion is not achieved within its applicable deadline after adding signed feature-extension days and excluding documented Client-delay days, the Provider shall refund only amounts actually received and allocated to the iOS Project. If no iOS amount has been received, no monetary refund is due and the Client has no obligation to pay the INR 50,000 iOS price. Any iOS refund due shall be completed within 30 calendar days after the missed iOS deadline.",
        "If any refund due under this Clause is not fully paid within the applicable 30-calendar-day refund period, the Provider shall pay the Client a delay penalty of INR 500 for each completed calendar day beginning immediately after that refund period expires and continuing until the full refund is paid. There is no maximum number of penalty days and no maximum monetary cap on this refund-delay penalty.",
        "Contract review and interpretation note. This note is intended to assist any human, legal or automated reviewer in understanding the Parties' commercial intent. The 30-calendar-day refund period is a negotiated payment-processing window, not a cancellation, waiver or postponement of the underlying refund obligation. It allows the Provider to reconcile amounts received, identify the correct project allocation, prepare any required invoice or tax adjustment and complete the transfer through normal banking channels. The Client remains protected because the refund obligation arises upon the applicable missed deadline, the full refund remains payable within that fixed period, and an uncapped delay penalty of INR 500 for each completed calendar day applies automatically after the period expires until payment is made. This explanatory note does not override or reduce any operative right or obligation in this Agreement.",
        "Failure or refund relating to one project does not cancel, reverse or make refundable the price for the other project if that other project was completed on time. Testing-track status is valid completion when Clauses 12 or 13 are satisfied."
    ])

    add_clause(doc, "16 Handover and Intellectual Property", [
        "Upon full payment of the Android Project price, the Client shall own the custom Android source code and project-specific shared backend and administration deliverables created for that project. Upon full payment of the iOS Project price, the Client shall own the custom iOS source code and iOS-specific deliverables.",
        "Each handover shall include the relevant source repository or archive, available build instructions, environment configuration, deployed-service details, setup notes and project-specific access. Delivery and payment for each project are intended to occur concurrently.",
        "The Provider warrants that project-specific custom source code delivered under this Agreement will be created by the Provider or its authorized team, or lawfully acquired and used. The Provider shall not knowingly include source code copied from another person's proprietary product without authorization. The Provider shall identify material third-party and open-source components in the handover information, and those components remain governed by their applicable licence terms.",
        "After payment for the applicable project, the Provider shall not sell or disclose a copy of the Client's paid project-specific custom source code to another client. The Provider retains its pre-existing tools, general know-how, independently developed concepts, reusable generic modules and third-party components while granting the Client the licence necessary to operate the delivered applications, subject to third-party licence terms.",
        "During development, the Provider shall maintain the project in a private version-controlled repository and make reasonable source-code commits or backups at least weekly while active coding is in progress. On written request, the Provider shall share a reasonable repository activity summary or progress evidence. Full repository access and source-code handover occur for each project in accordance with this Clause and the applicable payment obligation."
    ])

    add_clause(doc, "17 One Year Defect Maintenance and Delay Penalty", [
        "For one year after each project's respective Completion date, the Provider shall correct each reproducible defect in that project's delivered agreed functionality without an additional development charge. The Provider shall issue the corrected build or otherwise restore the agreed functionality within 24 hours after receiving a complete written defect report and all access reasonably required to reproduce and correct it.",
        "If a qualifying defect is not corrected within that 24-hour period, a delay penalty of INR 500 shall apply for each completed calendar day beginning immediately after the 24-hour period expires and continuing until the defect is corrected. There is no maximum number of penalty days and no maximum monetary cap on this delay penalty. Duplicate reports arising from the same root cause shall be treated as one defect. The accrued penalty may be adjusted against an unpaid project balance or, if the affected project has been fully paid, shall be paid to the Client within 30 calendar days after the Client's written demand.",
        "The 24-hour correction period and delay penalty do not apply while correction is blocked by missing Client information, credentials, access, approval or testing cooperation; to a new feature or changed requirement; or to an issue caused solely by Client misuse, an unsupported device or operating system, a third-party service outage, store review, network failure, external API change or platform-policy change. Any such blocked time is excluded after written notice to the Client. Ongoing infrastructure and third-party costs remain the Client's responsibility."
    ])

    add_clause(doc, "18 Termination and Exit", [
        "The Parties may terminate this Agreement, or either separate project, at any time by a written document signed by both Parties.",
        "Either Party may terminate an affected project for a material breach that the other Party does not cure within 15 calendar days after receiving complete written notice describing the breach and the action required to cure it. A Party may terminate immediately for fraud, deliberate misuse or unauthorized disclosure of confidential information, or insolvency where immediate termination is permitted by applicable law.",
        "The Client may terminate an affected project for convenience by giving 15 calendar days' written notice. The Provider shall give the Client a written statement of completed work and approved non-cancellable third-party commitments. The Client shall pay for completed work accepted up to the termination date and approved non-cancellable third-party costs. The Provider shall refund within 30 calendar days any amount received for the affected project that exceeds those sums. The token money shall be credited against the amount properly due and shall not be automatically forfeited. Clause 15 controls any refund caused by failure to meet a project deadline.",
        "The Provider may terminate an affected project if an undisputed payment default or Client-side blocking condition continues for 30 calendar days after written notice. The Provider shall deliver the work for which the Client has paid, subject to Clause 16. Termination of one project does not terminate the other unless the notice expressly states that both are affected. Confidentiality, intellectual-property rights, accrued payment and refund obligations, defect obligations for completed paid work and dispute provisions survive termination."
    ])

    add_clause(doc, "19 Client Responsibilities and Launch Readiness", [
        "The Client shall provide branding, content, policy documents, privacy policy, community rules, support contact, age or location rules, coin and gift values, game details, moderation decisions, Google Play Console access, Apple Developer and App Store Connect access, production server or cloud accounts, credentials, approved third-party accounts and timely review decisions.",
        "The Client should own production accounts so billing, data and credentials remain under its control. The Provider shall identify any missing input in writing. A Client-side blocking delay extends the affected deadline automatically under Clause 10; no separate signed extension is required for those excluded days.",
        "Store descriptions, screenshots, production account verification and store approvals are controlled or supplied by the Client or the relevant store. The Provider shall supply the agreed technical builds and submission information."
    ])

    add_clause(doc, "20 Confidentiality Data and Security", [
        "Each Party shall keep the other Party's non-public business information, credentials and project materials confidential and use them only for the project. The Provider shall apply reasonable access controls and shall not intentionally disclose Client or user data except as required to perform the services or by law.",
        "The Client is responsible for approving final privacy notices, consent language, age rules, retention decisions and community policies."
    ])

    add_clause(doc, "21 Governing Law and Courts", [
        "This Agreement is governed by the laws of India. Subject to applicable law, the competent courts at Bharuch, Gujarat shall have exclusive jurisdiction over disputes arising from or connected with this Agreement. Before filing proceedings, the Parties shall attempt in good faith to resolve the dispute through written discussion for 15 calendar days, unless urgent interim relief is required."
    ])

    add_clause(doc, "22 Notices Amendments and Electronic Signing", [
        "Formal notices and approvals shall be sent to the notified email addresses or verified messaging numbers. Amendments, accepted features, time revisions, price changes and waivers must be written and signed by both Parties.",
        "This Agreement may be signed in counterparts and by legally recognized electronic means. A scanned signed copy or valid electronic signature may be treated as an execution copy to the extent permitted by applicable law."
    ])

    add_clause(doc, "23 Entire Agreement Priority and Severability", [
        "This Agreement, Schedule A and signed feature addenda contain the complete agreement for Love Dost. A later signed addendum controls only the feature and additional time it expressly states; all remaining provisions of this Agreement continue unchanged.",
        "The Android and iOS project structure is severable. Invalidity, delay, refund or termination affecting one project does not automatically invalidate completed obligations for the other project. No earlier proposal, draft, chat or oral statement changes this Agreement unless incorporated into a document signed by both Parties."
    ])

    doc.add_heading("Schedule A Detailed Feature Acceptance and Account Responsibilities", level=1)
    doc.add_paragraph(
        "This Schedule states the agreed functional acceptance baseline for the Android Project and the iOS Project. A function passes when it performs the result stated below on the applicable representative devices using the Client-approved production or testing accounts. Differences required by Google, Apple or an approved third-party service do not constitute a defect if the agreed user result remains materially available."
    )
    schedule_items = [
        "A1 Android acceptance. The signed Android build shall install and open on the representative Android devices; complete mobile-number OTP login; create and edit a profile; discover users or rooms; join and operate agreed voice and video rooms; complete one-to-one voice and video calls; enter a configured paid private room; display and update the coin wallet; send an enabled gift; use the agreed YouTube watch-party controls; upload and play an eligible Reel; create one eligible Mere Vichar post; display ranks; submit a report; and reflect the relevant user, room, content, wallet and report data in administration. A blocking defect in any signed acceptance journey prevents Android acceptance until corrected.",
        "A2 iOS acceptance. The signed iOS build shall perform the same agreed user journeys as Android on the representative iPhones, subject only to Apple interface, permission, billing and policy requirements. The build shall connect to the agreed shared backend and administration system, preserve wallet and account data consistently and be uploaded to the Client's App Store Connect or TestFlight account. A blocking defect in any signed iOS acceptance journey prevents iOS acceptance until corrected.",
        "A3 Login and profile. The primary login method shall be mobile-number OTP using a Client-approved messaging service. The mandatory fields are a unique account identifier, display name and acceptance of the applicable terms and privacy notice. Profile photograph, introduction, gender, date of birth or age confirmation, language and location may be collected or displayed only when enabled by the Client's approved policy. OTP, SMS and identity-service charges are third-party costs payable by the Client.",
        "A4 Discovery and account safety. Users shall be able to view eligible profiles, live rooms and approved public content; search the enabled categories; block and unblock another user; report a user, room or content item; and update the available privacy and account settings. A blocked user shall not be able to initiate a direct interaction that the agreed blocking rule prohibits.",
        "A5 Voice and video party rooms. A host shall be able to create, open and close a room; control microphone, speaker and camera permissions; accept or reject microphone requests; remove a participant; and view the participant list, room chat and eligible gifts. At launch, a standard audio party room shall support at least 16 concurrent joined participant seats, including the host, and a standard video party room shall support at least 4 concurrent joined participant seats, including the host. These minimum party-room capacities are separate from the one-to-one call and the paid private-room limits in A6. The administration system shall enforce the configured speaker, camera-participant and audience limits. Capacity above these stated minimums and the available media quality remain subject to the Client-approved real-time service tier, device capability and network condition; load above the approved tier is not part of base acceptance.",
        "A6 One-to-one calls and paid private rooms. Two eligible users shall be able to start, receive, reject and end a private voice or video call. A paid private video room shall support no more than three joined participants. The entry coin price shall be visible before entry, sufficient balance shall be verified on the server, one successful entry shall create one deduction and one ledger record, and rejoining the same uninterrupted session shall not create a second charge. Failed entry shall not deduct coins. Interrupted-session adjustments may be made by an authorized administrator according to the recorded transaction.",
        "A7 Pricing, host earnings and settlement. Private-room prices, gift values, coin packs and any host earning percentage shall be configurable by an authorized administrator before launch. Rank points, coins, gifts or host earnings create no cash entitlement under the base scope. If the Client later requires real-money withdrawal or host payout, the Parties must first sign an addendum covering KYC, tax, payout-provider integration, fraud controls, settlement rules, vendor fees and additional time. Until that addendum is signed and implemented, no user or host has a contractual right to withdraw coins or earnings as cash.",
        "A8 YouTube watch party. A host shall be able to select a permitted embeddable YouTube video and use the approved play, pause and seek controls while room voice remains available. A late joiner shall receive the current available playback position. The app shall show a reasonable error when a video cannot be embedded or played. Availability remains subject to YouTube permissions, region, network and platform rules, and the application shall not sell access to the YouTube video.",
        "A9 Coins, gifts and billing. The administration system shall permit configuration of enabled coin packs and virtual gifts. The wallet shall show current balance and transaction history; server-side validation shall prevent a negative balance and duplicate successful deduction for the same transaction; and an eligible gift shall create the corresponding sender ledger entry and room event. Google Play Billing and Apple in-app purchase flows shall be used where platform rules require them. Store commissions, taxes, refunds and payment-provider charges remain governed by the applicable platform and are the Client's responsibility.",
        "A10 Games. The Provider shall install and integrate the Client-approved licensed third-party game SDK or module; the Provider is not developing original game source or content under the base scope. A clearly labelled Games button or menu shall appear inside each eligible voice or video party room in the room controls or bottom action toolbar, as finalized in the signed user-interface design. The control shall open the approved game catalogue, allow the host to launch an available game, allow eligible room participants to join and display the result or score supplied by the game service. If the approved vendor does not permit direct embedding, an agreed web view or deep-link launch method is acceptable if the agreed room-to-game user journey remains available. Game source code, artwork, licences, game-server charges and vendor availability are third-party matters. No particular game title is guaranteed until the Client approves an available licensed title and provides or pays for the required account, licence or SDK access.",
        "A11 Reels. An eligible user shall be able to upload an MP4 or MOV Reel of up to 60 seconds and 100 MB, subject to the approved storage and transcoding service. Vertical 9:16 media is the target presentation. The feed shall play eligible Reels and support the enabled like, comment and report actions; the uploader may delete the Reel; and an authorized administrator may review reports and disable or remove the Reel. Unsupported, corrupt or over-limit files may be rejected with an error message.",
        "A12 Mere Vichar. The daily posting limit shall use India Standard Time. Each eligible user may publish one text post of up to 500 characters per app calendar day. The user may edit or delete that post, and enabled reactions, comments, reports and administration moderation shall operate on it. Deleting a post does not permit a second post on the same day.",
        "A13 Ranks and leaderboard. Rank names, point values and eligible activities shall be configurable by an authorized administrator. The app shall display the user's current rank and an in-app leaderboard using the saved point total. Ties shall display the same point total and use earliest achievement time as the ordering tie-breaker. Repeated or suspicious events may be ignored, reversed or corrected by an administrator. The default leaderboard period is monthly unless the Client chooses a different available period before launch.",
        "A14 Administration and reports. The administration system shall provide Super Administrator and Moderator access levels. The Super Administrator may manage users, rooms, enabled content, gifts, coin settings, ranks, reports and operational configuration. A Moderator may review and act on assigned user, room and content reports without changing protected billing or ownership settings. Available audit and transaction records shall be retained for at least 90 days unless law, platform rules or the Client's approved privacy policy requires a different period.",
        "A15 Store and production accounts. Google Play Console, Apple Developer, App Store Connect and production server or cloud accounts shall be owned or controlled and funded by the Client. The Client shall provide the required account identifiers, agreements, tax and banking setup, verification, permissions and access by the date reasonably requested for the affected work and no later than ten calendar days before a planned store upload. Missing, expired or restricted access is a documented Client-delay period under Clause 10. Where technically suitable and available, the Client-approved production infrastructure shall use an India region as the preferred default in accordance with Clause 8.",
        "A16 Supported operating systems and devices. Final testing shall cover the current public major Android and iOS versions and the two immediately preceding major versions that remain supported by Google or Apple on the first formal acceptance-build date. Testing shall include at least one representative lower-capability Android phone, one current mainstream Android phone and two supported iPhone screen sizes. Device- or operating-system-specific behavior outside this matrix is not a blocking defect unless separately agreed.",
        "A17 Third-party services. The Provider may recommend suitable vendors for voice, video, messaging, OTP, storage, CDN, billing, analytics, crash reporting, games and other integrations, but activation requires the Client's written approval. Before approval, the Provider shall state the vendor, pricing basis, agreed usage assumptions and a reasonable cost estimate and shall record any approved spending cap. The Client shall own or fund the production account and pay the vendor's actual fees. Vendor outages, quotas, price or policy changes, unsupported regions and later API changes are not defects in the delivered application, although the Provider shall implement the agreed integration correctly.",
        "A18 Acceptance review. For each submitted project build, the Client shall test the signed journeys and provide one consolidated written defect list with steps, evidence and required access within seven calendar days. A blocking defect means a reproducible failure that prevents completion of a signed core journey with no reasonable workaround. Cosmetic differences, a new preference, changed content, an increased limit, a new game or integration, store-review timing or a third-party outage do not block Completion. Failure to report a blocking defect within the review period constitutes acceptance of the submitted build, without waiving the one-year defect obligations in Clause 17.",
        "A19 Formal contacts. Formal notices shall be sent to the Provider and Client email addresses or verified messaging numbers used for execution of this Agreement or later notified in writing. The Client's nominated reviewer is Samarjeet unless the Client appoints a replacement in writing. A change of contact does not change the scope, price or deadline.",
        "A20 Reserved future feature concepts. Lucky Bag is intended as an optional promotional feature through which an authorized administrator or eligible host may place a configured pool of virtual coins or virtual rewards into a time-limited bag for eligible room participants to claim. The intended implementation may use server-side randomized allocation, claim limits, ledger entries and anti-duplicate or anti-abuse checks. No user stake, cash withdrawal or cash entitlement is included under this preliminary description. PK War is intended as a timed competitive interaction between two eligible users, hosts or rooms, with pairing, start and end controls, live scores based on approved engagement or virtual-gift activity and a displayed result or winner. These descriptions record product direction only and are not final acceptance requirements. Before either feature becomes part of the agreed scope, the Parties must sign a Clause 11 feature addendum defining the applicable platforms, screens, participants, eligibility, funding source, coin accounting, claim or scoring rules, moderation, abuse controls, third-party dependencies, acceptance checks and exact additional calendar days. Until that addendum is signed, omission or non-operation of Lucky Bag or PK War is not a defect and does not prevent Android or iOS Completion."
    ]
    for item in schedule_items:
        doc.add_paragraph(item, style="List Bullet")
    add_signature_table(
        doc,
        "For the Provider",
        "For the Client",
        page_break=True,
        left_details={
            "name": "FORVILL ENTERPRISES PRIVATE LIMITED",
            "designation": "Authorized Signatory",
            "date": "02 October 2026",
            "place": "New Delhi, India",
        },
        right_details={
            "name": "Samarjeet",
            "designation": "Client",
            "date": "02 October 2026",
            "place": "Bharuch, Gujarat",
        },
    )
    CLIENT_APP_CONTRACT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(CLIENT_APP_CONTRACT_PATH)


if __name__ == "__main__":
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    build_main_english()
    print(CLIENT_APP_CONTRACT_PATH)
