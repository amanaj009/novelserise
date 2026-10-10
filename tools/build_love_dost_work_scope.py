from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from build_love_dost_contracts import (
    add_clause,
    add_page_number,
    add_title,
    configure_document,
    set_cell_margins,
    set_cell_shading,
    set_table_borders,
)


OUTPUT_DIR = Path(r"E:\AndroidStudioProjects\Love Dost App")
OUTPUT_PATH = OUTPUT_DIR / "Love-Dost-Work-Scope.docx"


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        for index, width in enumerate(widths):
            table.columns[index].width = Inches(width)
    for index, header in enumerate(headers):
        cell = table.cell(0, index)
        set_cell_shading(cell, "263B59")
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            if row_index % 2 == 0:
                set_cell_shading(cell, "F3F6FA")
            cell.paragraphs[0].add_run(value)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=110, start=120, bottom=110, end=120)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    set_table_borders(table)
    doc.add_paragraph()


def bullet_list(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build_scope():
    doc = Document()
    configure_document(doc, hindi=False)
    header = doc.sections[0].header.paragraphs[0]
    header.text = "FORVILL ENTERPRISE PVT. LTD.   |   LOVE DOST   |   WORK SCOPE"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor(75, 75, 75)

    add_title(
        doc,
        "LOVE DOST APPLICATION WORK SCOPE",
        "Android iOS Backend and Administration Delivery Plan",
        "Implementation reference for product development testing and handover",
    )

    add_clause(doc, "1 Scope Purpose", [
        "This document converts the Love Dost Software Development Agreement into an implementation checklist. It identifies the modules, platform ownership, dependencies, deliverables and completion checks that the project team must track. The signed Agreement remains the controlling commercial and legal document."
    ])

    add_clause(doc, "2 Project Structure", [
        "The work is organized as two independently completed applications supported by a shared backend and administration system. Android is developed natively in Kotlin and includes the initial shared backend and admin delivery. iOS is developed natively in Swift and integrates with the same approved services."
    ])
    add_table(
        doc,
        ["Workstream", "Technology", "Target", "Completion Evidence"],
        [
            ["Android application", "Kotlin", "100 Working Days", "Acceptance checks passed, Google Play Console testing deployment and source handover"],
            ["iOS application", "Swift", "120 Working Days", "Acceptance checks passed, App Store Connect or TestFlight upload and source handover"],
            ["Shared backend", "To be confirmed", "Delivered with Android foundation", "APIs, environment configuration, deployment details and project access"],
            ["Administration system", "To be confirmed", "Delivered with Android foundation", "Admin access and agreed user, room, content, wallet and moderation controls"],
        ],
        widths=[1.45, 1.2, 1.35, 2.75],
    )

    add_clause(doc, "3 Functional Work Packages")

    doc.add_heading("3.1 Accounts Profiles and Discovery", level=2)
    bullet_list(doc, [
        "Implement the approved registration and login method.",
        "Create profile setup and editing for display name, photograph, introduction and approved public fields.",
        "Implement user and live-room discovery plus agreed search or filter behavior.",
        "Add block, report, privacy and account-setting controls.",
    ])

    doc.add_heading("3.2 Live Voice and Video Party Rooms", level=2)
    bullet_list(doc, [
        "Create voice and video rooms with host, participant and audience roles.",
        "Implement room creation, joining, closing and member-list behavior.",
        "Implement microphone, speaker, camera and participation request controls.",
        "Implement agreed room chat, gift activity, reporting, removal and moderation.",
        "Validate approved room capacity and video-quality targets on representative devices and networks.",
    ])

    doc.add_heading("3.3 Private Calls and Paid Private Rooms", level=2)
    bullet_list(doc, [
        "Implement private one-to-one voice and video calls.",
        "Implement paid private video rooms for up to three joined participants.",
        "Show the coin price before entry, verify balance, deduct coins and write transaction history.",
        "Apply approved pricing authority, invite or discovery, rejoin and interrupted-session rules.",
    ])

    doc.add_heading("3.4 YouTube Watch Party", level=2)
    bullet_list(doc, [
        "Allow the host to select a permitted embeddable YouTube video.",
        "Implement approved synchronized play, pause, seek and late-join behavior.",
        "Keep room voice and moderation available during playback.",
        "Handle unavailable videos, host departure, region restrictions and player errors.",
    ])

    doc.add_heading("3.5 Wallet Coins Gifts and Games", level=2)
    bullet_list(doc, [
        "Implement coin balance, approved purchase flow and transaction history.",
        "Implement configurable gift catalogue, gift values and gift sending.",
        "Expose agreed gift and transaction controls to administrators.",
        "Implement only the game titles, assets, scoring and multiplayer rules approved in the signed feature list.",
    ])

    doc.add_heading("3.6 Reels and Mere Vichar", level=2)
    bullet_list(doc, [
        "Implement Reel upload, playback, vertical feed and approved like, comment and report behavior.",
        "Enforce approved duration, file size, format, storage and moderation limits.",
        "Allow one Mere Vichar written post per user per app calendar day.",
        "Implement approved timezone, text length, edit, delete, reaction, comment and reporting rules.",
    ])

    doc.add_heading("3.7 Ranks Leaderboard and Administration", level=2)
    bullet_list(doc, [
        "Implement user rank or level, points and leaderboard using approved scoring rules.",
        "Apply reset periods, ties and anti-abuse checks.",
        "Create administration controls for users, rooms, reports, gifts, coins, ranks, Reels, Mere Vichar and settings.",
        "Implement approved admin roles, moderation actions, audit information and retention behavior.",
    ])

    add_clause(doc, "4 Shared Backend and Technical Tasks", bullets=[
        "Define the approved backend architecture, data model, API contracts and environment structure.",
        "Implement authentication, authorization and access checks for user and admin functions.",
        "Implement room, call, wallet, transaction, gift, game, content, report and rank data flows.",
        "Implement media upload, processing, storage and delivery through approved third-party services.",
        "Add session and wallet-state handling, recoverable error messages and fault-investigation logging.",
        "Keep production accounts, billing and credentials under Client ownership wherever practical.",
        "Document environment variables, deployment configuration, service dependencies and backup expectations.",
    ])

    doc.add_page_break()
    add_clause(doc, "5 Delivery Sequence")
    add_table(
        doc,
        ["Stage", "Primary Work", "Review Evidence"],
        [
            ["Foundation", "Confirm requirements, Schedule A, screen flows, architecture, accounts and profiles", "Approved feature decisions, key screens and initial builds"],
            ["Live interaction", "Voice and video rooms, host controls, private calls and paid-room foundation", "Working Android and iOS demonstrations of agreed live journeys"],
            ["Social and monetization", "Watch party, wallet, gifts, games, Reels, Mere Vichar, ranks and admin", "Integrated review build with transaction and moderation flows"],
            ["Android release", "Android end-to-end QA, fixes, source and Google Play testing deployment", "Android Completion evidence"],
            ["iOS release", "iOS end-to-end QA, fixes, source and TestFlight or App Store Connect upload", "iOS Completion evidence"],
        ],
        widths=[1.25, 3.25, 2.25],
    )

    add_clause(doc, "6 Testing Checklist", bullets=[
        "Registration, login, logout, session recovery and profile editing.",
        "Room creation, discovery, joining, leaving, microphone and camera control.",
        "Private call success, failure, interruption and participant-limit behavior.",
        "Paid-room price display, insufficient balance, deduction and transaction records.",
        "Gift sending, coin history and administrator transaction review.",
        "YouTube playback synchronization and unavailable-video handling.",
        "Reel upload, validation, playback, feed behavior and reporting.",
        "Daily Mere Vichar restriction using the approved timezone.",
        "Rank calculation, leaderboard ordering, reset and anti-abuse behavior.",
        "Blocking, reporting, moderation and administrator permissions.",
        "Representative Android and iOS versions, devices and network conditions listed in the signed scope.",
    ])

    add_clause(doc, "7 Definition of Completion", [
        "Android is complete when its signed acceptance checks pass, it is deployed to the Client's Google Play Console on an internal, closed, open or production testing track, and the Android source plus agreed backend, admin, configuration and setup materials are available.",
        "iOS is complete when its signed acceptance checks pass, its release build and source are delivered, and it is uploaded to the Client's App Store Connect or TestFlight account. Public store approval is not required for either completion event."
    ])

    add_clause(doc, "8 Deliverables", bullets=[
        "Kotlin Android source code and release or testing build.",
        "Swift iOS source code and release or TestFlight build.",
        "Shared backend and administration source created for the project.",
        "Source repositories or archives and available build instructions.",
        "Environment configuration, deployment details and setup notes.",
        "Project-specific admin and service access created by the Provider.",
        "Known limitations and outstanding issue list, if any.",
    ])

    add_clause(doc, "9 Client Inputs and Third Party Dependencies", [
        "The Client must provide branding, content, legal policies, moderation decisions, account access, credentials, feature decisions and consolidated review feedback. All approved third-party costs are paid by the Client, including hosting, cloud, storage, voice or video services, SMS or OTP, payment gateway, Google Play, Apple Developer, APIs, purchased code and licences."
    ])

    add_clause(doc, "10 Change Control", [
        "An open or additional feature starts only after both Parties sign a short addendum identifying the affected platform, feature, dependencies, acceptance check and exact additional Working Days. An unsigned chat, call or demonstration does not change the scope or deadline. The controlling Agreement governs commercial treatment."
    ])

    add_clause(doc, "11 Predevelopment Decisions", bullets=[
        "Final Android and iOS screen lists and acceptance checks.",
        "Login method and profile fields.",
        "Room capacities, video targets and private-room rules.",
        "Coin, gift, settlement and payment behavior.",
        "Game titles, number, assets, scoring and multiplayer rules.",
        "Reel and Mere Vichar limits and moderation rules.",
        "Rank scoring, reset, ties and anti-abuse rules.",
        "Admin roles, reports, audit information and retention.",
        "Supported Android and iOS versions and representative devices.",
        "Approved voice, video, storage, messaging, billing and other service providers.",
    ])

    doc.add_heading("12 Scope Approval", level=1)
    add_table(
        doc,
        ["Provider", "Client"],
        [
            ["Name: __________________________", "Name: __________________________"],
            ["Signature: ______________________", "Signature: ______________________"],
            ["Date: ___________________________", "Date: ___________________________"],
        ],
        widths=[3.35, 3.35],
    )

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    build_scope()
